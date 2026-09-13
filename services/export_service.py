"""Синхронный экспорт заявок в Excel для Web API."""
from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from models import Registration, RegistrationStatus, Tournament, User
from utils.season import get_season_date_range, normalize_season_param, season_label


def fetch_registrations(
    session: Session,
    *,
    period: str = "all",
    month: Optional[str] = None,
    year: Optional[int] = None,
    season_key: Optional[str] = None,
) -> list[Registration]:
    query = session.query(Registration).join(Tournament).join(User)

    if period == "month" and month:
        query = query.filter(Tournament.month == month)
    elif period == "year" and year:
        start = datetime(year, 1, 1).date()
        end = datetime(year, 12, 31).date()
        query = query.filter(Tournament.date.between(start, end))
    elif period == "season":
        resolved = normalize_season_param(season_key)
        start, end = get_season_date_range(resolved)
        if start and end:
            query = query.filter(Tournament.date.between(start, end))
    elif season_key:
        resolved = normalize_season_param(season_key)
        start, end = get_season_date_range(resolved)
        if start and end:
            query = query.filter(Tournament.date.between(start, end))

    regs = query.all()
    regs.sort(key=lambda r: (r.tournament.date, r.tournament.name, r.user.last_name, r.user.first_name))
    return regs


def _period_title(period: str, month: Optional[str], year: Optional[int], season_key: Optional[str]) -> str:
    if period == "month" and month:
        return f"Месяц: {month}"
    if period == "year" and year:
        return f"Год: {year}"
    if period == "season":
        return f"Сезон: {season_label(normalize_season_param(season_key))}"
    if season_key:
        return f"Сезон: {season_label(normalize_season_param(season_key))}"
    return "Все сезоны"


def build_export_workbook_bytes(
    session: Session,
    *,
    period: str = "all",
    month: Optional[str] = None,
    year: Optional[int] = None,
    season_key: Optional[str] = None,
) -> tuple[Optional[BytesIO], str]:
    regs = fetch_registrations(session, period=period, month=month, year=year, season_key=season_key)
    title = _period_title(period, month, year, season_key)
    slug = (
        f"season_{normalize_season_param(season_key) or 'all'}"
        if period == "season"
        else f"month_{month}"
        if period == "month" and month
        else f"year_{year}"
        if period == "year" and year
        else "all"
    )
    filename = f"export_{slug}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

    if not regs:
        return None, filename

    wb = Workbook()
    ws = wb.active
    ws.title = "Заявки"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4472C4")

    ws.merge_cells("A1:F1")
    ws["A1"] = f"Отчёт по заявкам — {title}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Сформирован: {datetime.now().strftime('%d.%m.%Y %H:%M')}"

    headers = ["Дата", "Турнир", "Судья", "Функция", "Категория", "Статус"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    status_map = {
        RegistrationStatus.APPROVED: "Утверждено",
        RegistrationStatus.REJECTED: "Отклонено",
        RegistrationStatus.PENDING: "На рассмотрении",
    }

    row = 5
    for reg in regs:
        ws.cell(row=row, column=1, value=reg.tournament.date.strftime("%d.%m.%Y"))
        ws.cell(row=row, column=2, value=reg.tournament.name)
        ws.cell(row=row, column=3, value=f"{reg.user.first_name} {reg.user.last_name}")
        ws.cell(row=row, column=4, value=reg.user.function)
        ws.cell(row=row, column=5, value=reg.user.category)
        ws.cell(row=row, column=6, value=status_map.get(reg.status, reg.status))
        row += 1

    for idx, width in enumerate([14, 36, 24, 18, 14, 16], 1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    ws_stats = wb.create_sheet("Статистика судей")
    ws_stats.append(["Судья", "Всего", "Утверждено", "Отклонено", "На рассмотрении"])
    judge_stats: dict[str, dict[str, int]] = {}
    for reg in regs:
        name = f"{reg.user.first_name} {reg.user.last_name}"
        if name not in judge_stats:
            judge_stats[name] = {"total": 0, "approved": 0, "rejected": 0, "pending": 0}
        judge_stats[name]["total"] += 1
        if reg.status == RegistrationStatus.APPROVED:
            judge_stats[name]["approved"] += 1
        elif reg.status == RegistrationStatus.REJECTED:
            judge_stats[name]["rejected"] += 1
        else:
            judge_stats[name]["pending"] += 1

    for name, stats in sorted(judge_stats.items(), key=lambda x: (-x[1]["approved"], x[0])):
        ws_stats.append([
            name,
            stats["total"],
            stats["approved"],
            stats["rejected"],
            stats["pending"],
        ])

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer, filename
